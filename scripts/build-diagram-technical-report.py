#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build Word + PowerPoint technical report from diagrams.drawio (12 pages, expanded narrative)."""
from __future__ import annotations

import json
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))
from docs_paths import MD_OPERATIONS, TECH_PPT, TECH_WORD  # noqa: E402
from diagram_technical_report_content import (  # noqa: E402
    CLOSING,
    DATE_TH,
    EXEC_SUMMARY,
    INTRO,
    PAGE_COUNT,
    SECTIONS,
    VERSION,
)
from user_guide_process_context import (  # noqa: E402
    PPT_BODY_PT,
    PPT_FONT,
    PPT_NOTES_PT,
    PPT_TITLE_PT,
    WORD_BODY_PT,
    WORD_FONT,
    WORD_H1_PT,
    WORD_H2_PT,
    WORD_LINE_SPACING,
    WORD_TITLE_PT,
)

import importlib.util

_bpg_spec = importlib.util.spec_from_file_location(
    "build_portal_user_guides",
    Path(__file__).resolve().parent / "build-portal-user-guides.py",
)
_bpg = importlib.util.module_from_spec(_bpg_spec)
assert _bpg_spec.loader is not None
_bpg_spec.loader.exec_module(_bpg)
_add_heading = _bpg._add_heading
_add_para = _bpg._add_para
_add_toc = _bpg._add_toc
_add_word_table = _bpg._add_word_table
_apply_word_styles = _bpg._apply_word_styles
_ppt_fill_table = _bpg._ppt_fill_table
_ppt_set_font = _bpg._ppt_set_font
_set_run_font = _bpg._set_run_font

PAGES_DIR = REPO / "docs" / "diagrams" / "export" / "pages"
DOCX_OUT = TECH_WORD / "TECHNICAL_DIAGRAM_REPORT_TH.docx"
PPTX_OUT = TECH_PPT / "TECHNICAL_DIAGRAM_REPORT_PPT_TH.pptx"
MD_OUT = MD_OPERATIONS / "TECHNICAL_DIAGRAM_REPORT_TH.md"
IMAGE_WIDTH_IN = 6.5
PPT_IMAGE_WIDTH_IN = 11.8


def load_manifest() -> list[dict]:
    mf = PAGES_DIR / "manifest.json"
    if not mf.exists():
        raise FileNotFoundError(
            f"Missing {mf} — run: npm run diagrams:export"
        )
    data = json.loads(mf.read_text(encoding="utf-8"))
    pages = [p for p in data.get("pages", []) if not p.get("error")]
    pages = sorted(pages, key=lambda p: p.get("index", 0))[:PAGE_COUNT]
    if not pages:
        raise RuntimeError("No diagram pages exported")
    return pages


def add_centered_image(doc, img_path: Path, caption: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    if not img_path.exists():
        _add_para(doc, f"[ไม่พบไฟล์ภาพ: {img_path.name}]", bold=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(img_path), width=Inches(IMAGE_WIDTH_IN))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        _set_run_font(run, WORD_FONT, WORD_BODY_PT - 1)
    cap.paragraph_format.space_after = Pt(8)


def add_body_paragraphs(doc, paragraphs: list[str]) -> None:
    from docx.enum.text import WD_LINE_SPACING

    for text in paragraphs:
        p = doc.add_paragraph(text)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.paragraph_format.line_spacing = WORD_LINE_SPACING
        if p.runs:
            _set_run_font(p.runs[0], WORD_FONT, WORD_BODY_PT)


def add_bullets(doc, items: list[str]) -> None:
    from docx.enum.text import WD_LINE_SPACING

    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.paragraph_format.line_spacing = WORD_LINE_SPACING
        if p.runs:
            _set_run_font(p.runs[0], WORD_FONT, WORD_BODY_PT)


def sections_for_pages(page_count: int) -> list[dict]:
    return [s for s in SECTIONS if s["num"] <= page_count]


def build_word(pages: list[dict]) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    n = len(pages)
    sections = sections_for_pages(n)
    doc = Document()
    _apply_word_styles(doc)

    title_p = doc.add_paragraph("รายงานทางเทคนิค — สถาปัตยกรรม Izara Anywhere")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(title_p.runs[0], WORD_FONT, WORD_TITLE_PT, bold=True)
    sub = doc.add_paragraph(
        f"ฉบับอธิบายภาษาไทย {n} แผนภาพแรกจาก diagrams.drawio (v{VERSION}) "
        f"— ไฟล์แผนภาพอาจมีมากกว่า {n} แท็บ แต่รายงานนี้ไม่รวมแท็บถัดไป"
    )
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(sub.runs[0], WORD_FONT, WORD_H1_PT, bold=True)

    _add_word_table(
        doc,
        ["รายการ", "รายละเอียด"],
        [
            ["เวอร์ชัน", VERSION],
            ["วันที่จัดทำ", DATE_TH],
            ["แหล่งแผนภาพ", "docs/diagrams/diagrams.drawio"],
            ["จำนวนหน้าแผนภาพ", str(n)],
            ["Word รายงาน", str(DOCX_OUT.name)],
            ["PowerPoint", str(PPTX_OUT.name)],
            ["แบบอักษร", f"{WORD_FONT} {WORD_BODY_PT} pt"],
        ],
    )
    doc.add_page_break()
    _add_toc(doc)
    doc.add_page_break()

    _add_heading(doc, "บทนำ", level=1)
    for para in INTRO.split("\n\n"):
        _add_para(doc, para.strip())

    _add_heading(doc, "สรุปผู้บริหาร", level=2)
    _add_para(doc, EXEC_SUMMARY)

    file_by_index = {p["index"]: p for p in pages}

    for sec in sections:
        idx = sec["num"]
        pg = file_by_index.get(idx)
        img_path = PAGES_DIR / pg["file"] if pg else None

        doc.add_page_break()
        _add_heading(doc, f"บทที่ {idx} — {sec['title']}", level=1)
        _add_para(doc, sec["lead_in"])

        _add_heading(doc, "วัตถุประสงค์ของบทนี้", level=2)
        _add_para(doc, sec["purpose"])

        cap = f"ภาพที่ {idx}: {sec['diagram']} (diagrams.drawio หน้า {idx}/{n})"
        if img_path:
            add_centered_image(doc, img_path, cap)

        _add_heading(doc, "คำอธิบายและบริบท", level=2)
        add_body_paragraphs(doc, sec["body"])

        if sec.get("bullets"):
            _add_heading(doc, "ประเด็นสำคัญ", level=2)
            add_bullets(doc, sec["bullets"])

        if sec.get("table"):
            headers, rows = sec["table"]
            _add_word_table(doc, list(headers), [list(r) for r in rows])

        _add_heading(doc, "การอ่านแผนภาพ", level=2)
        _add_para(doc, sec["diagram_reading"])

        if sec.get("bridge"):
            _add_heading(doc, "เชื่อมต่อบทถัดไป", level=2)
            _add_para(doc, sec["bridge"])

    doc.add_page_break()
    _add_heading(doc, "บทสรุป", level=1)
    for para in CLOSING.split("\n\n"):
        _add_para(doc, para.strip())

    DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
    out_path = DOCX_OUT
    try:
        doc.save(str(out_path))
    except PermissionError:
        out_path = DOCX_OUT.with_stem(DOCX_OUT.stem + "_BUILD")
        doc.save(str(out_path))
        print(f"  WARN: {DOCX_OUT.name} is open — saved to {out_path.name}")
    print(f"  DOCX: {out_path} ({out_path.stat().st_size // 1024} KB)")


def build_ppt(pages: list[dict]) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    n = len(pages)
    sections = sections_for_pages(n)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    file_by_index = {p["index"]: p for p in pages}

    def slide_title(title: str, subtitle: str = "") -> None:
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(12.3), Inches(1.2))
        p = box.text_frame.paragraphs[0]
        p.text = title
        _ppt_set_font(p, PPT_TITLE_PT, bold=True)
        if subtitle:
            sub = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(12.3), Inches(0.6))
            sp = sub.text_frame.paragraphs[0]
            sp.text = subtitle
            _ppt_set_font(sp, PPT_BODY_PT)
        return slide

    slide_title(
        "รายงานทางเทคนิค — Izara Anywhere",
        f"v{VERSION} · {DATE_TH} · {n} แผนภาพ diagrams.drawio",
    )
    notes = prs.slides[-1].notes_slide.notes_text_frame
    notes.text = INTRO[:500]
    for para in notes.paragraphs:
        _ppt_set_font(para, PPT_NOTES_PT)

    ov = prs.slides.add_slide(blank)
    head = ov.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.7))
    head.text_frame.paragraphs[0].text = "สารบัญการนำเสนอ"
    _ppt_set_font(head.text_frame.paragraphs[0], PPT_TITLE_PT, bold=True)
    rows = [[str(s["num"]), s["title"]] for s in sections]
    tbl = ov.shapes.add_table(len(rows) + 1, 2, Inches(0.5), Inches(1.0), Inches(12), Inches(5.5)).table
    _ppt_fill_table(tbl, ["บท", "หัวข้อ"], rows, PPT_BODY_PT)

    for sec in sections:
        idx = sec["num"]
        pg = file_by_index.get(idx)
        img_path = PAGES_DIR / pg["file"] if pg else None

        slide = prs.slides.add_slide(blank)
        tbox = slide.shapes.add_textbox(Inches(0.35), Inches(0.12), Inches(12.6), Inches(0.65))
        tbox.text_frame.paragraphs[0].text = f"บทที่ {idx} — {sec['title']}"
        _ppt_set_font(tbox.text_frame.paragraphs[0], PPT_TITLE_PT, bold=True)

        if img_path and img_path.exists():
            slide.shapes.add_picture(
                str(img_path), Inches(0.35), Inches(0.85), width=Inches(PPT_IMAGE_WIDTH_IN)
            )
            body_top = 5.35
        else:
            body_top = 1.2

        bullets = [sec["lead_in"]] + sec.get("bullets", [])[:4]
        body = slide.shapes.add_textbox(Inches(0.4), Inches(body_top), Inches(12.3), Inches(1.9))
        tf = body.text_frame
        tf.word_wrap = True
        for bi, line in enumerate(bullets):
            para = tf.paragraphs[0] if bi == 0 else tf.add_paragraph()
            para.text = f"• {line}" if not line.startswith("•") else line
            _ppt_set_font(para, PPT_BODY_PT)

        notes_tf = slide.notes_slide.notes_text_frame
        note_lines = [sec["purpose"], "", *sec["body"][:3], "", sec.get("diagram_reading", "")]
        notes_tf.text = "\n".join(note_lines)
        for para in notes_tf.paragraphs:
            _ppt_set_font(para, PPT_NOTES_PT)

    slide_title("บทสรุป", "Izara Anywhere — สถาปัตยกรรมและ E2E")
    prs.slides[-1].notes_slide.notes_text_frame.text = CLOSING[:800]

    PPTX_OUT.parent.mkdir(parents=True, exist_ok=True)
    ppt_path = PPTX_OUT
    try:
        prs.save(str(ppt_path))
    except PermissionError:
        ppt_path = PPTX_OUT.with_stem(PPTX_OUT.stem + "_BUILD")
        prs.save(str(ppt_path))
        print(f"  WARN: {PPTX_OUT.name} is open — saved to {ppt_path.name}")
    print(f"  PPTX: {ppt_path} ({len(prs.slides)} slides, {ppt_path.stat().st_size // 1024} KB)")


def build_markdown(pages: list[dict]) -> None:
    n = len(pages)
    sections = sections_for_pages(n)
    lines = [
        "# รายงานทางเทคนิค — สถาปัตยกรรม Izara Anywhere",
        "",
        f"> v{VERSION} · {DATE_TH} · {n} แผนภาพ",
        "",
        "## บทนำ",
        "",
        INTRO,
        "",
        "## สรุปผู้บริหาร",
        "",
        EXEC_SUMMARY,
        "",
    ]
    file_by_index = {p["index"]: p for p in pages}
    for sec in sections:
        idx = sec["num"]
        pg = file_by_index.get(idx)
        rel = f"../diagrams/export/pages/{pg['file']}" if pg else ""
        lines += [f"## บทที่ {idx} — {sec['title']}", "", sec["lead_in"], ""]
        if rel:
            lines += [f"![{sec['diagram']}]({rel})", ""]
        lines += ["### วัตถุประสงค์", "", sec["purpose"], ""]
        for b in sec["body"]:
            lines += [b, ""]
        if sec.get("bridge"):
            lines += [sec["bridge"], ""]
    lines += ["## บทสรุป", "", CLOSING, ""]
    MD_OUT.write_text("\n".join(lines), encoding="utf-8")
    print(f"  MD:   {MD_OUT}")


def main() -> int:
    print(f"=== Build Diagram Technical Report ({PAGE_COUNT} pages target) ===")
    pages = load_manifest()
    print(f"  Using {len(pages)} diagram page(s)")
    build_word(pages)
    build_ppt(pages)
    build_markdown(pages)
    print("Done.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
