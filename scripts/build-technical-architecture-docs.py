#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build Thai technical architecture Word (.docx) and PowerPoint (.pptx)."""
from __future__ import annotations

import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))
from docs_paths import TECH_PPT, TECH_WORD  # noqa: E402

from technical_architecture_content import DATE_TH, SLIDES, VERSION  # noqa: E402
from user_guide_process_context import (  # noqa: E402
    PPT_BODY_PT,
    PPT_FONT,
    PPT_NOTES_PT,
    PPT_SUBTITLE_PT,
    PPT_TITLE_PT,
    WORD_BODY_PT,
    WORD_FONT,
    WORD_H1_PT,
    WORD_H2_PT,
    WORD_LINE_SPACING,
    WORD_TITLE_PT,
)

# Reuse Word helpers from portal guide builder (module name has hyphen)
import importlib.util

_bpg_spec = importlib.util.spec_from_file_location(
    "build_portal_user_guides",
    Path(__file__).resolve().parent / "build-portal-user-guides.py",
)
_bpg = importlib.util.module_from_spec(_bpg_spec)
assert _bpg_spec.loader is not None
_bpg_spec.loader.exec_module(_bpg)
_add_heading = _bpg._add_heading
_add_para = _bpg._add_para
_add_toc = _bpg._add_toc
_add_word_table = _bpg._add_word_table
_apply_word_styles = _bpg._apply_word_styles
_ppt_fill_table = _bpg._ppt_fill_table
_ppt_set_font = _bpg._ppt_set_font
_set_run_font = _bpg._set_run_font


def build_word(docx_path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Pt

    doc = Document()
    _apply_word_styles(doc)

    title_p = doc.add_paragraph("รายงานโครงสร้างทางเทคนิค — Izara Anywhere")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(title_p.runs[0], WORD_FONT, WORD_TITLE_PT, bold=True)
    sub = doc.add_paragraph("ระบบแพทย์ทางไกล (Telemedicine) — ฉบับภาษาไทย")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(sub.runs[0], WORD_FONT, WORD_H1_PT, bold=True)

    _add_word_table(
        doc,
        ["รายการ", "รายละเอียด"],
        [
            ["เวอร์ชัน", VERSION],
            ["วันที่จัดทำ", DATE_TH],
            ["แบบอักษรเนื้อหา", f"{WORD_FONT} {WORD_BODY_PT} pt"],
            ["ระยะบรรทัด", str(WORD_LINE_SPACING)],
            ["สไลด์นำเสนอ", f"{PPT_FONT} หัวข้อ {PPT_TITLE_PT} pt / เนื้อหา {PPT_BODY_PT} pt"],
            ["แผนภาพ", "docs/diagrams/diagrams.drawio (13 หน้า)"],
            ["ล้างข้อมูลทดสอบ", "npm run cleanup:cloud-test-only"],
        ],
    )
    doc.add_page_break()
    _add_toc(doc)
    doc.add_page_break()

    _add_heading(doc, "บทนำ", level=1)
    _add_para(
        doc,
        "เอกสารฉบับนี้อธิบายโครงสร้างทางเทคนิคของระบบ Izara Anywhere ครอบคลุมสถาปัตยกรรม "
        "เครือข่ายคลาวด์ ฐานข้อมูล การเชื่อมต่อบริการ และกระบวนการทำงานเป็นขั้นตอน "
        "จัดทำตามมาตรฐานการรายงานภาษาไทยของหน่วยงานราชการและสาธารณสุข "
        f"(แบบอักษร {WORD_FONT} ขนาดเนื้อหา {WORD_BODY_PT} pt ระยะบรรทัด {WORD_LINE_SPACING})",
    )

    for i, slide in enumerate(SLIDES, 1):
        _add_heading(doc, f"{i}. {slide['title']}", level=1)
        if slide.get("table"):
            headers, rows = slide["table"]
            _add_word_table(doc, list(headers), [list(r) for r in rows])
        for line in slide.get("body", []):
            p = doc.add_paragraph(line, style="List Bullet")
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            p.paragraph_format.line_spacing = WORD_LINE_SPACING
            if p.runs:
                _set_run_font(p.runs[0], WORD_FONT, WORD_BODY_PT)
        if slide.get("notes"):
            _add_para(doc, f"บันทึกวิทยากร: {slide['notes']}", size=WORD_BODY_PT)

    doc.save(str(docx_path))
    print(f"  DOCX: {docx_path} ({docx_path.stat().st_size // 1024} KB)")


def build_ppt(out_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for slide_def in SLIDES:
        slide = prs.slides.add_slide(blank)
        head = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.9))
        hp = head.text_frame.paragraphs[0]
        hp.text = slide_def["title"]
        _ppt_set_font(hp, PPT_TITLE_PT, bold=True)

        y = 1.15
        if slide_def.get("table"):
            headers, rows = slide_def["table"]
            n = min(len(rows), 8)
            tbl = slide.shapes.add_table(
                n + 1, len(headers), Inches(0.45), Inches(y), Inches(12.2), Inches(0.35 * (n + 2))
            ).table
            _ppt_fill_table(tbl, list(headers), [list(r) for r in rows[:n]], PPT_BODY_PT)
            y += 0.35 * (n + 2) + 0.15

        body = slide.shapes.add_textbox(Inches(0.45), Inches(y), Inches(12.2), Inches(7.5 - y - 0.3))
        tf = body.text_frame
        tf.word_wrap = True
        for idx, line in enumerate(slide_def.get("body", [])):
            para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            para.text = f"• {line}" if not line.startswith("•") else line
            _ppt_set_font(para, PPT_BODY_PT)

        if slide_def.get("notes"):
            notes_tf = slide.notes_slide.notes_text_frame
            notes_tf.text = slide_def["notes"]
            for para in notes_tf.paragraphs:
                _ppt_set_font(para, PPT_NOTES_PT)

    prs.save(str(out_path))
    print(f"  PPTX: {out_path} ({len(SLIDES)} slides, {PPT_FONT}, {out_path.stat().st_size // 1024} KB)")


def main() -> int:
    print("=== Build Technical Architecture docs (TH) ===")
    docx = TECH_WORD / "TECHNICAL_ARCHITECTURE_WORD_TH.docx"
    pptx = TECH_PPT / "TECHNICAL_ARCHITECTURE_PPT_TH.pptx"
    build_word(docx)
    build_ppt(pptx)
    print("Done.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
