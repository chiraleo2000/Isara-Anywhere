#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Build Thai user guides (Word .docx + PowerPoint .pptx) per portal from UI test screenshots.
"""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))
from docs_paths import DOCS, GUIDES_DOCTOR, GUIDES_PATIENT  # noqa: E402
from user_guide_catalog import (  # noqa: E402
    DATE_TH,
    DOCTOR_SHOTS,
    DOCTOR_URL,
    PATIENT_SHOTS,
    PATIENT_URL,
    VERSION,
    Shot,
)
from user_guide_process_context import (  # noqa: E402
    DOCTOR_PROCESS_INTRO,
    PATIENT_PROCESS_INTRO,
    PPT_BODY_PT,
    PPT_FONT,
    PPT_NOTES_PT,
    PPT_SUBTITLE_PT,
    PPT_TITLE_PT,
    WORD_BODY_PT,
    WORD_FONT,
    WORD_H1_PT,
    WORD_H2_PT,
    WORD_LINE_SPACING,
    WORD_TITLE_PT,
    section_context,
)

UI_GROUPS_PATIENT = "A, B, D, E, F, G, H, I, J, SSO, Workflows"
UI_GROUPS_DOCTOR = "A, C, D, E, F, G, H, I, SSO, Workflows (Doctor/Admin)"


def existing_shots(shots: list[Shot]) -> list[Shot]:
    out: list[Shot] = []
    for s in shots:
        p = REPO / s.path.replace("/", "\\") if "\\" in str(REPO) else REPO / s.path
        if p.is_file():
            out.append(s)
        else:
            print(f"  skip missing: {s.path}", file=sys.stderr)
    return out


def word_md(portal: str, shots: list[Shot], url: str, groups: str) -> str:
    title = (
        "พอร์ทัลผู้ป่วย (Patient Portal)"
        if portal == "patient"
        else "พอร์ทัลแพทย์และผู้ดูแลระบบ (Doctor & Admin Portal)"
    )
    lines = [
        f"# คู่มือการใช้งาน — {title}",
        "## ระบบ Isara Anywhere — เอกสารอธิบายการใช้งาน (ฉบับภาษาไทย)",
        "",
        "| รายละเอียดเอกสาร | ข้อมูล |",
        "|---|---|",
        "| ชื่อระบบ | Isara Anywhere — ระบบแพทย์ทางไกล |",
        f"| เวอร์ชัน | {VERSION} |",
        f"| URL (Cloud Dev-Testing) | {url} |",
        f"| วันที่จัดทำ | {DATE_TH} |",
        f"| กลุ่มทดสอบ UI | {groups} |",
        "| สถานะ | ใช้เฉพาะภาพจากการทดสอบ UI ที่ผ่าน (Passing Screenshots) |",
        f"| แบบอักษร Word | {WORD_FONT} {WORD_BODY_PT} pt (มาตรฐานรายงานภาษาไทย) |",
        f"| แบบอักษร PowerPoint | {PPT_FONT} |",
        "",
        "---",
        "",
        "## บทนำ",
        "",
        (PATIENT_PROCESS_INTRO if portal == "patient" else DOCTOR_PROCESS_INTRO).strip(),
        "",
    ]
    if portal == "patient":
        lines += [
            "**Patient Portal** ออกแบบสำหรับผู้ป่วย ครอบคลุมการจองนัดหมาย ติดตามสถานะ "
            "เข้าร่วมวิดีโอคอล บันทึก PHR ปรึกษา AI ค้นหาแพทย์และสถานพยาบาล "
            "รวมถึงการจัดการความยินยอม PDPA และหนังสือแสดงเจตนา",
            "",
            "### ข้อกำหนดเบื้องต้น",
            "",
            "| รายการ | ข้อกำหนด |",
            "|---|---|",
            "| เบราว์เซอร์ | Google Chrome หรือ Microsoft Edge เวอร์ชันล่าสุด |",
            "| อินเทอร์เน็ต | ≥ 5 Mbps (วิดีโอคอลแนะนำ 10 Mbps) |",
            "| อุปกรณ์ | กล้องและไมโครโฟนสำหรับ telehealth |",
            "| เอกสารอ้างอิง | `Processes/Appointment_Workflows.md`, `Processes/VIDEO_MEETING_JITSI_GEMINI.md` |",
            "",
        ]
    else:
        lines += [
            "**Doctor & Admin Portal** ใช้ URL เดียวกัน แต่แสดงเมนูตามบทบาท:",
            "",
            "| บทบาท | ขอบเขต | หน้าที่หลัก |",
            "|---|---|---|",
            "| **แพทย์ (Doctor)** | ผู้ป่วยของตนเอง | คิว, Health Meeting, EMR, วิดีโอคอล, เนื้อหา |",
            "| **ผู้ดูแล (Admin)** | ทั้งองค์กร | จัดสรรแพทย์, Pool, อนุมัติบัญชี/เนื้อหา |",
            "",
            "### เอกสารกระบวนการที่เกี่ยวข้อง",
            "",
            "- `Processes/Pages/Doctor-Portal/00_Doctor_Portal_Overview.md` — เส้นทางและเมนู",
            "- `Processes/Appointment_Workflows.md` — Pool, จัดสรร, ยืนยันนัด",
            "- `Processes/VIDEO_MEETING_JITSI_GEMINI.md` — Lobby, HOST, Guest",
            "",
        ]

    current_section = ""
    for i, s in enumerate(shots, 1):
        if s.section != current_section:
            current_section = s.section
            lines += ["", f"## {current_section}", ""]
            ctx = section_context(portal, current_section)
            if ctx:
                purpose, wf_steps, cautions, pref = ctx
                lines += [
                    "**วัตถุประสงค์ของหมวดนี้:**",
                    "",
                    purpose,
                    "",
                    "**ลำดับกระบวนการในระบบ (สรุป):**",
                    "",
                ]
                for ws in wf_steps:
                    lines.append(f"- {ws}")
                lines += ["", "**ข้อควรระวัง:**", ""]
                for c in cautions:
                    lines.append(f"- {c}")
                lines += ["", f"**อ้างอิง:** `{pref}`", ""]
        lines += [
            f"### {i}. {s.title}",
            "",
            "**วัตถุประสงค์ของขั้นตอน:**",
            "",
            s.script,
            "",
            "**ขั้นตอนการใช้งาน (ละเอียด):**",
            "",
        ]
        for j, step in enumerate(s.steps, 1):
            lines.append(f"{j}. {step}")
            lines.append(f"   - **ปฏิบัติ:** ดำเนินการตาม UI จนจบขั้นนี้ — อย่าข้าม modal ยืนยัน")
            lines.append(f"   - **ตรวจบนหน้าจอ:** {s.screen}")
            lines.append(f"   - **บทบาท:** {'ผู้ป่วย' if portal == 'patient' else 'แพทย์/ผู้ดูแล'}")
            lines.append(
                f"   - **รายละเอียด:** ขั้นที่ {j} ต้องสำเร็จก่อนขั้นถัดไป; บันทึก appointmentId/meetingId จาก Network หากต้องส่งต่อ IT"
            )
            lines.append("   - **ตรวจสอบ:** HTTP 2xx บน API หลัก; ไม่มี toast error")
        ctx = section_context(portal, s.section)
        if ctx:
            _, _, cautions, pref = ctx
            lines += [
                "",
                "**ข้อควรปฏิบัติในขั้นนี้:**",
                "",
            ]
            for c in cautions[:2]:
                lines.append(f"- {c}")
            lines.append(f"- อ้างอิงกระบวนการ: `{pref}`")
        lines += [
            "",
            "**ผลลัพธ์ที่คาดหวัง:**",
            "",
            f"- {s.screen}",
            "- ดำเนินการต่อได้โดยไม่มีข้อผิดพลาด (HTTP 4xx/5xx) บนฟังก์ชันหลัก",
            "- สอดคล้อง `tests/SELECTORS.md` และ Playwright group ที่เกี่ยวข้อง",
            "",
            "**คำอธิบายเพิ่มเติม / บริบทกระบวนการ:**",
            "",
            f"กลุ่มทดสอบ UI: `{s.section}` — ภาพนี้ยืนยันว่าขั้นตอน «{s.title}» ทำงานครบตามที่ออกแบบใน Processes/Pages "
            f"และผ่านการทดสอบอัตโนมัติบน Cloud dev-testing (v{VERSION})",
            f"เอกสาร Word ใช้ {WORD_FONT} {WORD_BODY_PT} pt; สไลด์ใช้ {PPT_FONT}.",
            "",
            f"**ลักษณะหน้าจอที่สำเร็จ:** {s.screen}",
            "",
            f"> **ภาพหน้าจอ:** `{s.path}`",
            f"> *{s.title} — จากการทดสอบ UI*",
            "",
            f"![{s.title}]({s.path})",
            "",
        ]

    lines += [
        "---",
        "",
        "## ภาคผนวก — สรุปกลุ่มทดสอบ UI",
        "",
        f"| จำนวนภาพในคู่มือ | {len(shots)} |",
        f"| กลุ่มทดสอบ | {groups} |",
        "",
        f"*จัดทำ {DATE_TH} — Isara Anywhere v{VERSION}*",
        "",
    ]
    return "\n".join(lines)


def ppt_md(portal: str, shots: list[Shot], url: str) -> str:
    title = (
        "คู่มือการใช้งาน — พอร์ทัลผู้ป่วย"
        if portal == "patient"
        else "คู่มือการใช้งาน — พอร์ทัลแพทย์และผู้ดูแลระบบ"
    )
    lines = [
        f"# สคริปต์นำเสนอสไลด์ — {title}",
        "## ระบบ Isara Anywhere — หนึ่งภาพต่อหนึ่งสไลด์ (ภาษาไทย)",
        "",
        f"> URL: {url} | เวอร์ชัน {VERSION} | {DATE_TH}",
        "",
    ]
    slide = 1

    # Slide 1 — Title
    lines += [
        f"## สไลด์ที่ {slide}",
        "",
        "**ชื่อสไลด์:** ชื่อเรื่อง — Isara Anywhere",
        "",
        "**สคริปต์บรรยาย:**",
        f'> "สวัสดีครับ/ค่ะ วันนี้นำเสนอคู่มือการใช้งาน {title} '
        f'เวอร์ชัน {VERSION} ทุกภาพหน้าจอมาจากการทดสอบ UI ที่ผ่านทั้งหมด"',
        "",
        "---",
        "",
    ]
    slide += 1

    # Slide 2 — Overview table
    if portal == "patient":
        overview_script = (
            "Patient Portal มี 10 เมนูหลัก ครอบคลุมจองนัด PHR วิดีโอคอล AI "
            "แผนที่ PDPA และคลังความรู้"
        )
        overview_body = (
            "| หัวข้อ | รายละเอียด |\n|---|---|\n"
            f"| URL | {url} |\n"
            "| บทบาท | ผู้ป่วย |\n"
            "| เมนูหลัก | 10 รายการ |\n"
            f"| กลุ่มทดสอบ | {UI_GROUPS_PATIENT} |"
        )
    else:
        overview_script = (
            "Portal เดียวรองรับแพทย์และผู้ดูแล แพทย์ดูเฉพาะผู้ป่วยของตน "
            "ผู้ดูแลจัดสรรแพทย์และอนุมัติระบบ"
        )
        overview_body = (
            "| หัวข้อ | แพทย์ | ผู้ดูแล |\n|---|---|---|\n"
            f"| URL | {url} | เดียวกัน |\n"
            "| Pool | รับเอง | จัดสรรให้แพทย์ |\n"
            f"| กลุ่มทดสอบ | {UI_GROUPS_DOCTOR} |"
        )

    lines += [
        f"## สไลด์ที่ {slide}",
        "",
        "**ชื่อสไลด์:** ตารางภาพรวมระบบ",
        "",
        overview_body,
        "",
        "**สคริปต์บรรยาย:**",
        f'> "{overview_script}"',
        "",
        "---",
        "",
    ]
    slide += 1

    for s in shots:
        steps_text = "\n".join(f"{n}. {st}" for n, st in enumerate(s.steps, 1))
        lines += [
            f"## สไลด์ที่ {slide}",
            "",
            f"**ชื่อสไลด์:** {s.title}",
            "",
            f"**กลุ่ม:** {s.section}",
            "",
            f"[ภาพ: {s.path}]",
            "",
            "**วัตถุประสงค์:**",
            "",
            s.script,
            "",
            "**ขั้นตอนบนสไลด์:**",
            "",
            steps_text,
            "",
            "**รายละเอียดเพิ่มเติม (บันทึกวิทยากร — FC Iconic 16 pt):**",
            "",
            f"- วัตถุประสงค์: {s.script}",
            f"- ตรวจบนหน้าจอ: {s.screen}",
            "- ปฏิบัติทีละขั้น — อย่าข้าม modal ยืนยัน",
            "- บันทึก appointmentId/meetingId จาก URL หรือ DevTools เมื่อเกี่ยวข้อง",
            "- ตรวจ HTTP 2xx และไม่มี toast แดงก่อนไปขั้นถัดไป",
            "- อ้างอิง: Processes/Pages + `tests/SELECTORS.md` + `Processes/Appointment_Workflows.md`",
            "",
            "**สคริปต์บรรยาย:**",
            f'> "ต่อไปคือ {s.title} ในกลุ่ม {s.section}. {s.script} '
            f'ขั้นตอนคือ: {" ".join(s.steps)} '
            f'เมื่อทำสำเร็จบนหน้าจอจะเห็น{s.screen} '
            f'เอกสาร Word ใช้ {WORD_FONT} {WORD_BODY_PT} pt สไลด์ใช้ {PPT_FONT} '
            f'อ้างอิงกระบวนการใน Processes/Pages และ UI test ที่ผ่านแล้ว"',
            "",
            "---",
            "",
        ]
        slide += 1

    # Closing slide
    lines += [
        f"## สไลด์ที่ {slide}",
        "",
        "**ชื่อสไลด์:** สรุปและ Q&A",
        "",
        "**สคริปต์บรรยาย:**",
        f'> "สรุปคู่มือ {title} ครบ {len(shots)} ขั้นตอนพร้อมภาพจาก UI tests '
        f'ขอบคุณครับ/ค่ะ มีคำถามเพิ่มเติมไหมครับ/คะ"',
        "",
    ]
    return "\n".join(lines)


def inject_word_images(md: str) -> str:
    pattern = re.compile(
        r"(> \*\*ภาพหน้าจอ:\*\* `([^`]+)`\n> \*([^\*]+)\*)\n(?!\!\[)",
        re.MULTILINE,
    )

    def repl(m: re.Match) -> str:
        path, caption = m.group(2), m.group(3)
        return f"{m.group(1)}\n\n![{caption}]({path})\n"

    return pattern.sub(repl, md)


def _set_run_font(run, font_name: str, size_pt: int, bold: bool = False) -> None:
    from docx.shared import Pt

    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    try:
        from docx.oxml.ns import qn

        r = run._element
        rPr = r.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:eastAsia"), font_name)
    except Exception:
        pass


def _add_para(doc, text: str, *, font=WORD_FONT, size=WORD_BODY_PT, bold=False, style=None):
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Pt

    p = doc.add_paragraph(text, style=style)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = WORD_LINE_SPACING
    p.paragraph_format.space_after = Pt(6)
    if p.runs:
        _set_run_font(p.runs[0], font, size, bold)
    return p


def _apply_word_styles(doc) -> None:
    from docx.shared import Pt

    normal = doc.styles["Normal"]
    normal.font.name = WORD_FONT
    normal.font.size = Pt(WORD_BODY_PT)
    for level, size in ((1, WORD_H1_PT), (2, WORD_H2_PT)):
        st = doc.styles[f"Heading {level}"]
        st.font.name = WORD_FONT
        st.font.size = Pt(size)
        st.font.bold = True


def _add_heading(doc, text: str, level: int = 1):
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Pt

    p = doc.add_paragraph(text, style=f"Heading {level}")
    size = WORD_H1_PT if level == 1 else WORD_H2_PT
    for run in p.runs:
        _set_run_font(run, WORD_FONT, size, bold=True)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = WORD_LINE_SPACING
    return p


def _add_word_table(doc, headers: list[str], rows: list[list[str]]) -> None:
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Pt

    if not rows:
        return
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ci, header in enumerate(headers):
        cell = table.rows[0].cells[ci]
        cell.text = header
        for para in cell.paragraphs:
            para.paragraph_format.space_after = Pt(2)
            for run in para.runs:
                _set_run_font(run, WORD_FONT, WORD_BODY_PT, bold=True)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                for run in para.runs:
                    _set_run_font(run, WORD_FONT, WORD_BODY_PT)
    doc.add_paragraph()


def _add_toc(doc) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    _add_heading(doc, "สารบัญ", level=1)
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)
    _add_para(
        doc,
        "หมายเหตุ: เปิดใน Microsoft Word แล้วคลิกขวาที่สารบัญ → «อัปเดตฟิลด์» (หรือกด F9) เพื่อแสดงหมายเลขหน้า",
        size=WORD_BODY_PT,
    )


def _role_label(portal: str) -> str:
    return "ผู้ป่วย" if portal == "patient" else "แพทย์/ผู้ดูแล"


def _section_pref(portal: str, section: str) -> str:
    ctx = section_context(portal, section)
    return ctx[3] if ctx else "Processes/Pages/"


def _enriched_steps(s: Shot) -> list[str]:
    """Expand each catalog step with verification sub-steps for Word/PPT tables."""
    out: list[str] = []
    for j, step in enumerate(s.steps, 1):
        out.append(step)
        out.append(f"ตรวจหน้าจอขั้นที่ {j}: {s.screen}")
        if j == len(s.steps):
            out.append("ตรวจ Network: HTTP 2xx, ไม่มี toast แดง, session/JWT ยังใช้ได้")
            out.append("บันทึก appointmentId / meetingId จาก URL หรือ DevTools หากต้องส่งต่อ IT")
    return out


def _step_table_rows(s: Shot, portal: str, pref: str) -> list[list[str]]:
    role = _role_label(portal)
    rows: list[list[str]] = []
    enriched = _enriched_steps(s)
    for j, step in enumerate(enriched, 1):
        is_check = step.startswith("ตรวจ") or step.startswith("บันทึก")
        rows.append(
            [
                str(j),
                step,
                "ดำเนินการจนจบ — อย่าข้าม modal ยืนยัน" if not is_check else "ตรวจสอบตามข้อความ",
                s.screen if not is_check else "HTTP 2xx / ไม่มี error banner",
                role,
                pref if j == 1 else "",
            ]
        )
    return rows


def build_docx_python(
    portal: str,
    shots: list[Shot],
    url: str,
    groups: str,
    docx_path: Path,
    doc_title: str,
) -> bool:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt
    except ImportError:
        print("  python-docx not installed — falling back to pandoc", file=sys.stderr)
        return False

    doc = Document()
    _apply_word_styles(doc)

    title_p = doc.add_paragraph(doc_title)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(title_p.runs[0], WORD_FONT, WORD_TITLE_PT, bold=True)
    sub = doc.add_paragraph("ระบบ Isara Anywhere — คู่มือการใช้งานฉบับภาษาไทย")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(sub.runs[0], WORD_FONT, WORD_H1_PT, bold=True)

    _add_word_table(
        doc,
        ["รายการ", "รายละเอียด"],
        [
            ["เวอร์ชันระบบ", VERSION],
            ["วันที่จัดทำ", DATE_TH],
            ["URL (Cloud Dev-Testing)", url],
            ["กลุ่มทดสอบ UI", groups],
            ["จำนวนขั้นตอนในคู่มือ", str(len(shots))],
            ["มาตรฐานรายงาน", f"{WORD_FONT} {WORD_BODY_PT} pt, ระยะบรรทัด {WORD_LINE_SPACING}"],
        ],
    )
    _add_word_table(
        doc,
        ["ประเภทเอกสาร", "แบบอักษร", "ขนาดตัวอักษร"],
        [
            ["Word / รายงาน PDF", WORD_FONT, f"{WORD_BODY_PT} pt (เนื้อหา)"],
            ["หัวข้อระดับ 1", WORD_FONT, f"{WORD_H1_PT} pt"],
            ["หัวข้อระดับ 2", WORD_FONT, f"{WORD_H2_PT} pt (ตัวหนา)"],
            ["ชื่อเรื่อง", WORD_FONT, f"{WORD_TITLE_PT} pt"],
            ["PowerPoint", PPT_FONT, f"หัวข้อ {PPT_TITLE_PT} pt / เนื้อหา {PPT_BODY_PT} pt"],
        ],
    )
    doc.add_page_break()
    _add_toc(doc)
    doc.add_page_break()

    _add_heading(doc, "บทนำ", level=1)
    intro = PATIENT_PROCESS_INTRO if portal == "patient" else DOCTOR_PROCESS_INTRO
    for block in intro.strip().split("\n"):
        line = block.strip()
        if line and not line.startswith("|") and not line.startswith("-"):
            _add_para(doc, line)

    current_section = ""
    role = _role_label(portal)
    for i, s in enumerate(shots, 1):
        if s.section != current_section:
            current_section = s.section
            doc.add_page_break()
            _add_heading(doc, current_section, level=1)
            ctx = section_context(portal, current_section)
            if ctx:
                purpose, wf_steps, cautions, pref = ctx
                _add_para(doc, "วัตถุประสงค์ของหมวด", bold=True)
                _add_para(doc, purpose)
                _add_word_table(
                    doc,
                    ["ลำดับ", "ขั้นตอนกระบวนการในระบบ"],
                    [[str(n), ws] for n, ws in enumerate(wf_steps, 1)],
                )
                _add_word_table(
                    doc,
                    ["ข้อควรระวัง", "รายละเอียด"],
                    [[str(n), c] for n, c in enumerate(cautions, 1)],
                )
                _add_para(doc, f"อ้างอิงกระบวนการ: {pref}", bold=True)

        _add_heading(doc, f"{i}. {s.title}", level=2)
        _add_word_table(
            doc,
            ["หัวข้อ", "รายละเอียด"],
            [
                ["กลุ่ม UI", s.section],
                ["บทบาทผู้ใช้", role],
                ["วัตถุประสงค์", s.script],
                ["ภาพหน้าจอ", s.path],
            ],
        )
        pref = _section_pref(portal, s.section)
        _add_word_table(
            doc,
            ["ลำดับ", "ขั้นตอน / รายละเอียด", "การปฏิบัติ", "การตรวจสอบ", "บทบาท", "อ้างอิง Processes"],
            _step_table_rows(s, portal, pref),
        )
        ctx = section_context(portal, s.section)
        cautions = ctx[2] if ctx else []
        _add_word_table(
            doc,
            ["ผลลัพธ์ที่คาดหวัง", "รายละเอียด"],
            [
                ["หน้าจอ", s.screen],
                ["API", "ไม่มี HTTP 4xx/5xx บนฟังก์ชันหลัก"],
                ["ทดสอบ", "สอดคล้อง tests/SELECTORS.md และ Playwright"],
                ["กระบวนการ", f"อ้างอิง {pref}"],
            ],
        )
        if cautions:
            _add_word_table(
                doc,
                ["ข้อควรปฏิบัติ", "คำอธิบาย"],
                [[str(n), c] for n, c in enumerate(cautions[:3], 1)],
            )
        _add_para(
            doc,
            f"คำอธิบายเพิ่มเติม: ขั้นตอน «{s.title}» ผ่านการทดสอบ UI บน Cloud dev-testing "
            f"(v{VERSION}) — จัดทำตามมาตรฐานรายงานภาษาไทย ({WORD_FONT} {WORD_BODY_PT} pt)",
        )
        img = REPO / s.path.replace("/", "\\") if "\\" in str(REPO) else REPO / s.path
        if img.is_file():
            try:
                doc.add_picture(str(img), width=Inches(6.2))
            except Exception as ex:
                _add_para(doc, f"(ไม่สามารถแนบภาพ: {ex})")
        doc.add_paragraph()

    doc.add_page_break()
    _add_heading(doc, "ภาคผนวก — สรุป", level=1)
    _add_word_table(
        doc,
        ["รายการ", "ค่า"],
        [
            ["จำนวนภาพในคู่มือ", str(len(shots))],
            ["กลุ่มทดสอบ", groups],
            ["เวอร์ชัน", VERSION],
            ["วันที่", DATE_TH],
        ],
    )

    tmp_docx = docx_path.with_suffix(".docx.building")
    doc.save(str(tmp_docx))
    try:
        if docx_path.exists():
            docx_path.unlink()
        tmp_docx.replace(docx_path)
    except OSError as ex:
        alt = docx_path.with_name(docx_path.stem + "_NEW.docx")
        tmp_docx.replace(alt)
        print(
            f"  DOCX: {alt} ({alt.stat().st_size // 1024} KB) — close {docx_path.name} and rename",
            file=sys.stderr,
        )
        print(f"  (could not overwrite {docx_path}: {ex})", file=sys.stderr)
        return True
    print(
        f"  DOCX: {docx_path} ({docx_path.stat().st_size // 1024} KB, {WORD_FONT} {WORD_BODY_PT}pt, TOC+tables)"
    )
    return True


def build_docx(md_path: Path, docx_path: Path, doc_title: str, portal: str, shots: list[Shot], url: str, groups: str) -> bool:
    if build_docx_python(portal, shots, url, groups, docx_path, doc_title):
        return True
    raw = md_path.read_text(encoding="utf-8")
    enriched = inject_word_images(raw)
    tmp = md_path.parent / f"_pandoc_{md_path.name}"
    tmp.write_text(enriched, encoding="utf-8")
    cmd = [
        "pandoc",
        str(tmp),
        "-o",
        str(docx_path),
        "--resource-path",
        str(REPO),
        "-s",
        "--toc",
        "--toc-depth=3",
        "--metadata",
        "lang=th",
        "--metadata",
        f"title={doc_title}",
    ]
    try:
        subprocess.run(cmd, check=True, capture_output=True, text=True)
        print(f"  DOCX (pandoc fallback): {docx_path} ({docx_path.stat().st_size // 1024} KB)")
        return True
    except (FileNotFoundError, subprocess.CalledProcessError) as e:
        print(f"  DOCX failed: {e}", file=sys.stderr)
        return False
    finally:
        tmp.unlink(missing_ok=True)


def parse_ppt_slides(md: str) -> list[dict]:
    slides: list[dict] = []
    blocks = re.split(r"(?=## สไลด์ที่ \d+)", md)
    for block in blocks:
        m_num = re.search(r"## สไลด์ที่ (\d+)", block)
        if not m_num:
            continue
        num = int(m_num.group(1))
        m_title = re.search(r"\*\*ชื่อสไลด์:\*\* (.+)", block)
        title = m_title.group(1).strip() if m_title else f"สไลด์ {num}"
        m_img = re.search(r"\[ภาพ:\s*([^\]\n]+)\]", block)
        image = None
        if m_img:
            raw = m_img.group(1).strip()
            image = REPO / raw.replace("/", "\\") if "\\" in str(REPO) else REPO / raw
        m_notes = re.search(
            r"\*\*สคริปต์บรรยาย:\*\*\s*\n>\s*\"(.+?)\"",
            block,
            re.DOTALL,
        )
        notes = m_notes.group(1).strip() if m_notes else ""
        steps_lines: list[str] = []
        m_steps = re.search(
            r"\*\*ขั้นตอนบนสไลด์:\*\*\s*\n((?:\d+\. .+\n?)+)",
            block,
        )
        if m_steps:
            steps_lines = [
                ln.strip()
                for ln in m_steps.group(1).strip().splitlines()
                if ln.strip()
            ]
        slides.append(
            {
                "num": num,
                "title": title,
                "image": image,
                "notes": notes,
                "steps_lines": steps_lines,
            }
        )
    return sorted(slides, key=lambda s: s["num"])


def _ppt_set_font(paragraph, size_pt: int, bold: bool = False) -> None:
    from pptx.util import Pt

    paragraph.font.name = PPT_FONT
    paragraph.font.size = Pt(size_pt)
    paragraph.font.bold = bold


def _ppt_fill_table(table, headers: list[str], rows: list[list[str]], font_pt: int) -> None:
    for ci, header in enumerate(headers):
        cell = table.cell(0, ci)
        cell.text = header
        _ppt_set_font(cell.text_frame.paragraphs[0], font_pt, bold=True)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri + 1, ci)
            cell.text = str(val)[:200]
            _ppt_set_font(cell.text_frame.paragraphs[0], max(font_pt - 2, 12))


def build_pptx_from_shots(
    shots: list[Shot],
    out_path: Path,
    portal: str,
    url: str,
    doc_title: str,
) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    role = _role_label(portal)
    subtitle = "Patient Portal" if portal == "patient" else "Doctor & Admin Portal"
    groups = UI_GROUPS_PATIENT if portal == "patient" else UI_GROUPS_DOCTOR

    def add_title_slide(title: str, lines: list[str], notes: str = "") -> None:
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.3), Inches(1.2))
        p = box.text_frame.paragraphs[0]
        p.text = title
        _ppt_set_font(p, PPT_TITLE_PT, bold=True)
        body = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(12), Inches(5.2))
        tf = body.text_frame
        tf.word_wrap = True
        for idx, line in enumerate(lines):
            para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            para.text = line
            _ppt_set_font(para, PPT_BODY_PT)
        if notes:
            notes_tf = slide.notes_slide.notes_text_frame
            notes_tf.text = notes
            for para in notes_tf.paragraphs:
                _ppt_set_font(para, PPT_NOTES_PT)

    add_title_slide(
        doc_title,
        [
            f"Isara Anywhere — {subtitle}",
            f"v{VERSION} | {DATE_TH}",
            f"แบบอักษร {PPT_FONT} | มาตรฐานสไลด์ภาษาไทย",
            url,
        ],
        f"นำเสนอคู่มือ {doc_title} เวอร์ชัน {VERSION}",
    )

    overview = prs.slides.add_slide(blank)
    tbox = overview.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.8))
    tbox.text_frame.paragraphs[0].text = "ภาพรวมระบบและมาตรฐานเอกสาร"
    _ppt_set_font(tbox.text_frame.paragraphs[0], PPT_SUBTITLE_PT, bold=True)
    tbl = overview.shapes.add_table(6, 2, Inches(0.5), Inches(1.1), Inches(12.2), Inches(2.8)).table
    _ppt_fill_table(
        tbl,
        ["รายการ", "รายละเอียด"],
        [
            ["URL", url],
            ["กลุ่มทดสอบ", groups],
            ["จำนวนสไลด์เนื้อหา", str(len(shots))],
            ["Word", f"{WORD_FONT} {WORD_BODY_PT} pt"],
            ["PowerPoint", PPT_FONT],
        ],
        PPT_BODY_PT,
    )

    current_section = ""
    for i, s in enumerate(shots, 1):
        if s.section != current_section:
            current_section = s.section
            ctx = section_context(portal, s.section)
            if ctx:
                purpose, wf_steps, cautions, pref = ctx
                sec_lines = [purpose] + [f"{n}. {w}" for n, w in enumerate(wf_steps, 1)]
                add_title_slide(
                    f"หมวด: {s.section}",
                    sec_lines[:6],
                    " ".join(cautions) + f" อ้างอิง {pref}",
                )

        slide = prs.slides.add_slide(blank)
        head = slide.shapes.add_textbox(Inches(0.35), Inches(0.12), Inches(12.6), Inches(0.75))
        hp = head.text_frame.paragraphs[0]
        hp.text = f"{i}. {s.title}"
        _ppt_set_font(hp, PPT_SUBTITLE_PT, bold=True)

        sub = slide.shapes.add_textbox(Inches(0.35), Inches(0.85), Inches(12.6), Inches(0.45))
        sp = sub.text_frame.paragraphs[0]
        sp.text = f"กลุ่ม {s.section} | {role} | วัตถุประสงค์: {s.script[:120]}"
        _ppt_set_font(sp, PPT_BODY_PT - 2)

        img_path = REPO / s.path.replace("/", "\\") if "\\" in str(REPO) else REPO / s.path
        if img_path.is_file():
            slide.shapes.add_picture(str(img_path), Inches(0.35), Inches(1.35), width=Inches(6.8))

        pref = _section_pref(portal, s.section)
        enriched = _enriched_steps(s)
        step_rows = [
            [str(j), st[:95], (s.screen[:35] if not st.startswith("ตรวจ") else "ตรวจ API/UI")[:35]]
            for j, st in enumerate(enriched, 1)
        ]
        if not step_rows:
            step_rows = [["1", s.script[:95], s.screen[:35]]]
        rows_n = min(len(step_rows), 10)
        font_pt = PPT_BODY_PT - 4 if rows_n > 7 else PPT_BODY_PT - 2
        stbl = slide.shapes.add_table(
            rows_n + 1,
            3,
            Inches(7.2),
            Inches(1.3),
            Inches(5.85),
            Inches(min(3.6, 0.32 * (rows_n + 1))),
        ).table
        _ppt_fill_table(
            stbl,
            ["ลำดับ", "ขั้นตอน (ละเอียด)", "ตรวจสอบ"],
            step_rows[:rows_n],
            font_pt,
        )
        if len(enriched) > rows_n:
            extra = slide.shapes.add_textbox(Inches(7.2), Inches(5.0), Inches(5.85), Inches(0.55))
            ep = extra.text_frame.paragraphs[0]
            ep.text = f"(ขั้นตอน {rows_n + 1}–{len(enriched)} ดูบันทึกวิทยากร)"
            _ppt_set_font(ep, font_pt)

        foot = slide.shapes.add_textbox(Inches(0.35), Inches(4.7), Inches(12.6), Inches(2.5))
        ft = foot.text_frame
        ft.word_wrap = True
        detail_lines = [
            f"ผลลัพธ์: {s.screen}",
            f"วัตถุประสงค์: {s.script}",
            "ตรวจ: HTTP 2xx, JWT/session, ไม่มี toast error",
            f"อ้างอิง: {pref}",
            f"ภาพ UI test: {s.path}",
        ]
        ctx = section_context(portal, s.section)
        if ctx:
            for c in ctx[2][:2]:
                detail_lines.append(f"ข้อควรระวัง: {c}")
        for idx, line in enumerate(detail_lines):
            para = ft.paragraphs[0] if idx == 0 else ft.add_paragraph()
            para.text = line
            _ppt_set_font(para, PPT_BODY_PT - 2)

        notes_parts = [
            f"สไลด์ {i}: {s.title} | กลุ่ม {s.section}",
            s.script,
            "ขั้นตอนละเอียด:",
            *[f"  {n}. {st}" for n, st in enumerate(enriched, 1)],
            f"ผลลัพธ์: {s.screen}",
            f"อ้างอิง {pref}",
            f"Word: {WORD_FONT} {WORD_BODY_PT} pt (ตาราง+สารบัญ) | สไลด์: {PPT_FONT}",
        ]
        notes_tf = slide.notes_slide.notes_text_frame
        notes_tf.text = "\n".join(notes_parts)
        for para in notes_tf.paragraphs:
            _ppt_set_font(para, PPT_NOTES_PT)

    add_title_slide(
        "สรุปและ Q&A",
        [
            f"ครบ {len(shots)} ขั้นตอนพร้อมภาพ UI tests",
            f"{WORD_FONT} {WORD_BODY_PT} pt (Word) / {PPT_FONT} (สไลด์)",
            "อ้างอิง Processes/Pages และ GATE0",
        ],
        "ขอบคุณ — เปิด Q&A",
    )

    tmp_pptx = out_path.with_suffix(".pptx.building")
    prs.save(str(tmp_pptx))
    try:
        if out_path.exists():
            out_path.unlink()
        tmp_pptx.replace(out_path)
    except OSError as ex:
        alt = out_path.with_name(out_path.stem + "_NEW.pptx")
        tmp_pptx.replace(alt)
        print(f"  PPTX: {alt} — close {out_path.name} and rename ({ex})", file=sys.stderr)
        out_path = alt
    img_count = sum(
        1
        for s in shots
        if (REPO / s.path.replace("/", "\\") if "\\" in str(REPO) else REPO / s.path).is_file()
    )
    print(
        f"  PPTX: {out_path} ({len(prs.slides)} slides, {img_count} images, "
        f"{PPT_FONT}, tables per slide, {out_path.stat().st_size // 1024} KB)"
    )


def build_pptx(slides: list[dict], out_path: Path, portal: str) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    subtitle = "Patient Portal" if portal == "patient" else "Doctor & Admin Portal"

    for slide_def in slides:
        slide = prs.slides.add_slide(blank)
        title_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.15), Inches(12.5), Inches(0.9))
        p = title_box.text_frame.paragraphs[0]
        p.text = slide_def["title"]
        _ppt_set_font(p, PPT_SUBTITLE_PT if len(slide_def["title"]) > 42 else PPT_TITLE_PT, bold=True)

        img_path = slide_def.get("image")
        has_image = img_path and Path(img_path).is_file()
        steps_lines = slide_def.get("steps_lines") or []

        if has_image:
            slide.shapes.add_picture(str(img_path), Inches(0.45), Inches(1.0), width=Inches(12.4))
            if steps_lines:
                foot = slide.shapes.add_textbox(Inches(0.45), Inches(6.55), Inches(12.4), Inches(0.75))
                tf = foot.text_frame
                tf.word_wrap = True
                for j, line in enumerate(steps_lines[:4]):
                    para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                    para.text = line
                    _ppt_set_font(para, PPT_BODY_PT - 2)
        elif slide_def["num"] == 1:
            sub = slide.shapes.add_textbox(Inches(0.6), Inches(2.0), Inches(12), Inches(3.5))
            t = sub.text_frame.paragraphs[0]
            t.text = f"Isara Anywhere\n{subtitle}\nv{VERSION} | {DATE_TH}\nแบบอักษร {PPT_FONT}"
            _ppt_set_font(t, PPT_TITLE_PT, bold=True)
        elif slide_def["num"] == 2:
            body = slide.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(12), Inches(5.5))
            url = PATIENT_URL if portal == "patient" else DOCTOR_URL
            for j, line in enumerate(
                [
                    f"URL: {url}",
                    f"สไลด์มีภาพจาก UI tests: {len(slides) - 3} รายการ",
                    "รูปแบบ: หนึ่งภาพต่อหนึ่งสไลด์ + คำอธิบายขั้นตอน",
                    "อ้างอิง Processes/ และ GATE0 workflows",
                ]
            ):
                para = body.text_frame.paragraphs[0] if j == 0 else body.text_frame.add_paragraph()
                para.text = line
                _ppt_set_font(para, PPT_BODY_PT)
        elif steps_lines:
            body = slide.shapes.add_textbox(Inches(0.6), Inches(1.2), Inches(12), Inches(5.8))
            for j, line in enumerate(steps_lines):
                para = body.text_frame.paragraphs[0] if j == 0 else body.text_frame.add_paragraph()
                para.text = line
                _ppt_set_font(para, PPT_BODY_PT)

        if slide_def.get("notes"):
            notes_tf = slide.notes_slide.notes_text_frame
            notes_tf.text = slide_def["notes"]
            for para in notes_tf.paragraphs:
                _ppt_set_font(para, PPT_NOTES_PT)

    prs.save(str(out_path))
    img_count = sum(1 for s in slides if s.get("image") and Path(s["image"]).is_file())
    print(
        f"  PPTX: {out_path} ({len(slides)} slides, {img_count} images, "
        f"{PPT_FONT}, {out_path.stat().st_size // 1024} KB)"
    )


def build_portal(portal: str) -> None:
    if portal == "patient":
        shots = existing_shots(PATIENT_SHOTS)
        url = PATIENT_URL
        groups = UI_GROUPS_PATIENT
        word_out = GUIDES_PATIENT / "USER_GUIDE_PATIENT_WORD_TH.md"
        docx_out = GUIDES_PATIENT / "USER_GUIDE_PATIENT_WORD_TH.docx"
        ppt_md_out = GUIDES_PATIENT / "USER_GUIDE_PATIENT_PPT_TH.md"
        pptx_out = GUIDES_PATIENT / "USER_GUIDE_PATIENT_PPT_TH.pptx"
        doc_title = "คู่มือผู้ป่วย Isara Anywhere"
    else:
        shots = existing_shots(DOCTOR_SHOTS)
        url = DOCTOR_URL
        groups = UI_GROUPS_DOCTOR
        word_out = GUIDES_DOCTOR / "USER_GUIDE_DOCTOR_WORD_TH.md"
        docx_out = GUIDES_DOCTOR / "USER_GUIDE_DOCTOR_WORD_TH.docx"
        ppt_md_out = GUIDES_DOCTOR / "USER_GUIDE_DOCTOR_PPT_TH.md"
        pptx_out = GUIDES_DOCTOR / "USER_GUIDE_DOCTOR_PPT_TH.pptx"
        doc_title = "คู่มือแพทย์และผู้ดูแล Isara Anywhere"

    print(f"\n=== {portal.upper()} ({len(shots)} screenshots) ===")

    # Intermediate .md is gitignored; deliverables are .docx / .pptx only
    word_out.write_text(word_md(portal, shots, url, groups), encoding="utf-8")
    ppt_md_out.write_text(ppt_md(portal, shots, url), encoding="utf-8")

    build_docx(word_out, docx_out, doc_title, portal, shots, url, groups)
    build_pptx_from_shots(shots, pptx_out, portal, url, doc_title)


def main() -> int:
    print("=== Build Isara Anywhere Portal User Guides (TH) ===")
    build_portal("patient")
    build_portal("doctor")
    print("\nDone. Output: docs/guides/patient|doctor/, docs/technical/")
    return 0


if __name__ == "__main__":
    sys.exit(main())
