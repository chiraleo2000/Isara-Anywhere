#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generate Thai user guides — per-portal Word (.docx) and PowerPoint (.pptx).

Preferred entry point:
  python scripts/build-portal-user-guides.py

This script still supports the legacy combined report (USER_GUIDE_WORD_REPORT_TH).
"""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
DOCS = REPO / "docs"
WORD_MD = DOCS / "USER_GUIDE_WORD_REPORT_TH.md"
WORD_DOCX = DOCS / "USER_GUIDE_WORD_REPORT_TH.docx"
PPT_MD = DOCS / "USER_GUIDE_PPT_SCRIPT_TH.md"
PPT_PPTX = DOCS / "USER_GUIDE_PPT_SCRIPT_TH.pptx"
BUILD_PORTAL = REPO / "scripts" / "build-portal-user-guides.py"


def inject_word_images(md: str) -> str:
    """Convert blockquote screenshot refs to markdown images for pandoc."""
    pattern = re.compile(
        r"(> \*\*ภาพหน้าจอ:\*\* `([^`]+)`\n> \*([^\*]+)\*)\n(?!\!\[)",
        re.MULTILINE,
    )

    def repl(m: re.Match) -> str:
        path, caption = m.group(2), m.group(3)
        if f"![{caption}]" in md[m.end() : m.end() + 200]:
            return m.group(0)
        return f"{m.group(1)}\n\n![{caption}]({path})\n"

    return pattern.sub(repl, md)


def build_docx_pandoc(md_path: Path, out_path: Path) -> bool:
    raw = md_path.read_text(encoding="utf-8")
    enriched = inject_word_images(raw)
    tmp = DOCS / "_USER_GUIDE_WORD_REPORT_TH_pandoc.md"
    tmp.write_text(enriched, encoding="utf-8")
    cmd = [
        "pandoc",
        str(tmp),
        "-o",
        str(out_path),
        "--resource-path",
        str(REPO),
        "-s",
        "--toc",
        "--toc-depth=3",
        "--metadata",
        "lang=th",
        "--metadata",
        "title=คู่มือการใช้งานระบบ Isara Anywhere",
    ]
    try:
        subprocess.run(cmd, check=True, capture_output=True, text=True)
        print(f"  DOCX: {out_path} ({out_path.stat().st_size // 1024} KB)")
        return True
    except FileNotFoundError:
        print("  pandoc not found — skipping DOCX", file=sys.stderr)
        return False
    except subprocess.CalledProcessError as e:
        print(f"  pandoc error: {e.stderr}", file=sys.stderr)
        return False
    finally:
        if tmp.exists():
            tmp.unlink(missing_ok=True)


def parse_ppt_slides(md: str) -> list[dict]:
    slides: list[dict] = []
    blocks = re.split(r"(?=## สไลด์ที่ \d+)", md)
    for block in blocks:
        m_num = re.search(r"## สไลด์ที่ (\d+)", block)
        if not m_num:
            continue
        num = int(m_num.group(1))
        m_title = re.search(r"\*\*ชื่อสไลด์:\*\* (.+)", block)
        title = m_title.group(1).strip() if m_title else f"สไลด์ {num}"
        m_img = re.search(r"\[ภาพ:\s*([^\]\n]+)\]", block)
        image = None
        if m_img:
            raw = m_img.group(1).strip()
            if raw.startswith("docs/"):
                image = REPO / raw.replace("/", "\\") if "\\" in str(REPO) else REPO / raw
            elif "screenshots" in raw:
                image = REPO / raw
        m_notes = re.search(
            r"\*\*สคริปต์บรรยาย:\*\*\s*\n>\s*\"(.+?)\"",
            block,
            re.DOTALL,
        )
        notes = m_notes.group(1).strip() if m_notes else ""
        slides.append({"num": num, "title": title, "image": image, "notes": notes})
    return sorted(slides, key=lambda s: s["num"])


def build_pptx(slides: list[dict], out_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for i, slide_def in enumerate(slides):
        slide = prs.slides.add_slide(blank)
        title_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.25), Inches(12.5), Inches(0.9))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = slide_def["title"]
        p.font.size = Pt(28)
        p.font.bold = True

        img_path = slide_def.get("image")
        has_image = img_path and Path(img_path).is_file()
        if has_image:
            slide.shapes.add_picture(
                str(img_path),
                Inches(0.5),
                Inches(1.2),
                width=Inches(12.3),
            )
        elif slide_def["num"] == 2:
            # Overview table slide
            body = slide.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(12), Inches(5))
            lines = [
                "พอร์ทัลผู้ป่วย — จองนัด, PHR, วิดีโอคอล (10 เมนู)",
                "พอร์ทัลแพทย์ — คิว, Health Meeting, Lobby, Jitsi (8 เมนู)",
                "พอร์ทัลผู้ดูแล — จัดสรรแพทย์, Pool, รายงานคิว",
                "GCP Dev-Testing: izara-*-dev-testing-724889190329.asia-southeast1.run.app",
            ]
            for j, line in enumerate(lines):
                para = body.text_frame.paragraphs[0] if j == 0 else body.text_frame.add_paragraph()
                para.text = line
                para.font.size = Pt(20)
        elif slide_def["num"] in (1, 46, 48):
            sub = slide.shapes.add_textbox(Inches(0.6), Inches(2), Inches(12), Inches(4))
            sub.text_frame.paragraphs[0].text = (
                "Isara Anywhere — ระบบแพทย์ทางไกล\n"
                "เวอร์ชัน 1.7.3 | พฤษภาคม 2569\n"
                "ผู้ป่วย | แพทย์ | ผู้ดูแลระบบ"
            )
            sub.text_frame.paragraphs[0].font.size = Pt(22)
        else:
            hint = slide.shapes.add_textbox(Inches(0.6), Inches(3), Inches(12), Inches(1))
            hint.text_frame.paragraphs[0].text = "(ไม่มีภาพหน้าจอจากการทดสอบสำหรับสไลด์นี้)"
            hint.text_frame.paragraphs[0].font.size = Pt(18)

        if slide_def.get("notes"):
            notes_slide = slide.notes_slide
            notes_slide.notes_text_frame.text = slide_def["notes"]

    prs.save(str(out_path))
    print(f"  PPTX: {out_path} ({len(slides)} slides, {out_path.stat().st_size // 1024} KB)")


def main() -> int:
    print("=== Generate Isara Anywhere Thai User Guides ===\n")
    if BUILD_PORTAL.is_file():
        print("Building per-portal guides (Patient + Doctor/Admin)...")
        subprocess.run([sys.executable, str(BUILD_PORTAL)], check=False)
        print()
    if not WORD_MD.is_file():
        print(f"Missing {WORD_MD}", file=sys.stderr)
        return 1
    if not PPT_MD.is_file():
        print(f"Missing {PPT_MD}", file=sys.stderr)
        return 1

    print("Building Word document...")
    ok_docx = build_docx_pandoc(WORD_MD, WORD_DOCX)

    print("Building PowerPoint...")
    ppt_md = PPT_MD.read_text(encoding="utf-8")
    slides = parse_ppt_slides(ppt_md)
    if not slides:
        print("  No slides parsed", file=sys.stderr)
        return 1
    build_pptx(slides, PPT_PPTX)

    print("\nDone.")
    return 0 if ok_docx else 0


if __name__ == "__main__":
    sys.exit(main())
